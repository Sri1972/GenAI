"""
Excel Creation Agent — designs a workbook plan with Claude (via the Claude
Code CLI, on Bedrock) and builds it with openpyxl.

Usage:
    python run.py --brief "Q3 sales by region and product" [--context "..."] [--out out.xlsx]
"""

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from common.claude_cli import run_claude, extract_json

SYSTEM_PROMPT = """You are the Excel Creation Agent. You have no tool access
in this call -- the brief and any reference content below are already
complete. Do not ask to open files or run scripts yourself; design the plan
directly from what's given. Given a content brief (and optionally reference
content from an upstream source), design a workbook of one or more sheets
with real data/rows, not placeholders.

Respond with ONLY a single JSON object, no prose, matching this shape:
{
  "title": "<workbook title, used as the filename>",
  "sheets": [
    {"name": "<sheet name>", "headers": ["col1", "col2", "..."],
     "rows": [["value1", "value2", "..."]]}
  ]
}
Rules:
- Every row must have the same number of values as there are headers.
- Use actual numbers (not strings) for numeric columns where appropriate.
- Never fabricate specific facts/figures that aren't in the brief or
  reference content -- if exact numbers aren't given, build a clearly
  labeled structure (headers, category labels) rather than inventing false
  precision."""


def build_xlsx(plan: dict, out_path: Path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_plan in plan.get("sheets", []):
        ws = wb.create_sheet(title=(sheet_plan.get("name") or "Sheet")[:31])
        headers = sheet_plan.get("headers", [])
        if headers:
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
        for row in sheet_plan.get("rows", []):
            ws.append(row)
        for col_cells in ws.columns:
            max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 40)

    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.save(str(out_path))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--brief", required=True, help="What the workbook should contain")
    ap.add_argument("--context", help="Optional reference content (e.g. from an upstream source) to base the workbook on")
    ap.add_argument("--out")
    args = ap.parse_args()

    prompt_parts = [f"BRIEF: {args.brief}"]
    if args.context:
        prompt_parts.append(f"\nREFERENCE CONTENT:\n{args.context}")

    result = run_claude("\n".join(prompt_parts), system_prompt=SYSTEM_PROMPT)
    plan = extract_json(result["result"])

    out_path = Path(args.out) if args.out else Path.cwd() / f"{(plan.get('title') or 'workbook').strip().replace(' ', '_')[:60]}.xlsx"
    build_xlsx(plan, out_path)

    n_rows = sum(len(s.get("rows", [])) for s in plan.get("sheets", []))
    print(f"Wrote {len(plan.get('sheets', []))} sheet(s), {n_rows} data row(s) to {out_path}")


if __name__ == "__main__":
    main()
