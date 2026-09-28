"""
PPT Creation Agent — designs a slide-by-slide deck plan with Claude (via the
Claude Code CLI, on Bedrock) and builds it with python-pptx, using real
native PowerPoint objects (charts, tables, chevron process rows, Gantt-style
timelines) rather than just title+bullets+picture. Three template modes:

  --template deck.pptx   Fill an existing template: Claude picks from that
                          template's real slide layouts/placeholders.
  --template doc.pdf      Reference only (a PDF has no editable slide
                          structure) — Claude designs a new deck "in the
                          spirit of" the PDF's structure/images, built from
                          the default template below.
  (no --template)         Falls back to DEFAULT_TEMPLATE_PATH (the org's own
                          MobilityGlobal template) if present, else
                          python-pptx's generic built-in layouts.

Usage:
    python run.py --brief "Q3 board update: revenue, churn, roadmap" [--template deck.pptx] [--out out.pptx] [--slides 8]

Native-feature notes (see generate()/build_pptx() below):
  - Charts and tables are fully native, editable-in-PowerPoint objects.
  - Chevrons are real MSO_SHAPE.CHEVRON AutoShapes in a row -- NOT SmartArt;
    python-pptx has no API to create SmartArt graphics at all.
  - Gantt is a row of positioned ROUNDED_RECTANGLE bars -- PowerPoint has no
    native Gantt object, so this (real editable shapes on a computed
    timeline scale) is the standard technique any tool uses for this.
  - Icons come from ICON_CATALOG below (32 PNGs extracted from the org
    template's own "Icons" style-guide slide, bundled in ./icons/) -- NOT
    PowerPoint's own native Icons gallery, which has no creation API either.
"""

import argparse
import json
import sys
import tempfile
from pathlib import Path

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from common.claude_cli import run_claude, extract_json

# The org's standard template -- used whenever a request doesn't upload its
# own --template, so PPT Creator is on-brand by default rather than
# python-pptx's generic blank layouts. Bundled here rather than referencing
# a personal machine path; only the 46 slide LAYOUTS matter at generation
# time (never the example content slides a template ships with), so the
# smaller abbreviated copy is exactly as useful as the full one.
DEFAULT_TEMPLATE_PATH = Path(__file__).resolve().parent / "templates" / "mobilityglobal_default.pptx"

# Pulled from that template's own ppt/theme/theme1.xml <a:clrScheme> --
# reused for chevron/Gantt fills so those native shapes stay on-brand. Only
# accurate for this specific template; a different uploaded template won't
# automatically match these (out of scope for now -- see run.py's own notes
# above about staying within native-object scope, not general theme
# extraction for arbitrary templates).
THEME_ACCENTS = ["0064D2", "420E71", "132445", "B8EAF5", "E3ABFF", "FFE783"]

CHART_KIND_TO_XL = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
    "area": XL_CHART_TYPE.AREA,
}

# 32 icons extracted directly from the org template's own "Icons" style-
# guide slide (bundled here as plain PNGs, ./icons/icon_NN.png) -- named by
# hand from a visual review, since the source slide has no per-icon labels.
# NOT PowerPoint's own native Icons gallery (Insert > Icons) -- that has no
# creation API at all, this is a separate, self-contained asset set. Line
# art is a fixed dark-navy stroke baked into the pixels, so these are only
# legible on a "...Light" background -- see SYSTEM_PROMPT's icon rules.
ICONS_DIR = Path(__file__).resolve().parent / "icons"
ICON_CATALOG = {
    "arrow": "icon_00.png", "notification": "icon_01.png", "guide": "icon_02.png",
    "bookmark": "icon_03.png", "briefcase": "icon_04.png", "calendar": "icon_05.png",
    "vehicle": "icon_06.png", "checkmark": "icon_07.png", "chart": "icon_08.png",
    "document": "icon_09.png", "download": "icon_10.png", "folder": "icon_11.png",
    "milestone": "icon_12.png", "favorite": "icon_13.png", "home": "icon_14.png",
    "inbox": "icon_15.png", "speed": "icon_16.png", "email": "icon_17.png",
    "map": "icon_18.png", "chat": "icon_19.png", "send": "icon_20.png",
    "person": "icon_21.png", "location": "icon_22.png", "dashboard": "icon_23.png",
    "print": "icon_24.png", "search": "icon_25.png", "security": "icon_26.png",
    "performance": "icon_27.png", "pricing": "icon_28.png", "delete": "icon_29.png",
    "settings": "icon_30.png", "close": "icon_31.png",
}

# Shared AWS/GCP/Azure/general icon catalog (see ContentAgents/icons/) --
# "aws"/"gcp"/"azure" are curated from each cloud provider's own official
# icon package; "general" is a small set of non-cloud business/automotive
# concept icons (dealership, car, owner, supplier, OEM) from Tabler Icons --
# neither is this template's own icons. Loaded here as name -> resolved PNG
# path (python-pptx's insert_picture has no SVG support, so the PNG side of
# that shared catalog is what this agent uses -- see icons/build_catalog.py
# for how both sides are produced). Despite the variable's name (kept for
# minimal diff -- it predates "general"), this covers all of those
# categories, not just cloud providers. Selected via "<category>:<name>"
# (e.g. "aws:ec2", "general:dealership"), validated against this dict
# alongside the flat ICON_CATALOG above in _resolve_icon_path -- kept as a
# SEPARATE dict, never merged into ICON_CATALOG, so a plain name always
# means this template's own icon and a "category:name" name always means a
# real cloud/business glyph from the shared catalog.
SHARED_ICONS_DIR = Path(__file__).resolve().parent.parent / "icons"
try:
    _CLOUD_CATALOG_RAW = json.loads((SHARED_ICONS_DIR / "catalog.json").read_text(encoding="utf-8"))
except (FileNotFoundError, json.JSONDecodeError):
    _CLOUD_CATALOG_RAW = {}
CLOUD_ICON_CATALOG = {
    provider: {name: SHARED_ICONS_DIR / provider / "png" / f"{stem}.png" for name, stem in services.items()}
    for provider, services in _CLOUD_CATALOG_RAW.items()
}


def _resolve_icon_path(icon_name: str) -> Path | None:
    """Resolves an "icons" item's "icon" field to a real file -- either a
    plain name from this template's own ICON_CATALOG, or a "<provider>:
    <service>" name from CLOUD_ICON_CATALOG (e.g. "aws:ec2"). Returns None
    for anything that doesn't resolve, same permissive-drop posture as
    every other catalog lookup in this file."""
    if ":" in icon_name:
        provider, service = icon_name.split(":", 1)
        return CLOUD_ICON_CATALOG.get(provider, {}).get(service)
    filename = ICON_CATALOG.get(icon_name)
    return (ICONS_DIR / filename) if filename else None


def _readable_text_color(hex_color: str) -> str:
    """Same relative-luminance rule as the Visualization Agent's own JS
    textColorOn() helper (see visualization_agent/run.py) -- three of
    THEME_ACCENTS are light pastels (B8EAF5/E3ABFF/FFE783), not just the
    three dark ones, so hardcoding white text on "whichever accent color
    cycles round" made some chevron steps render white-on-light and
    unreadable. Picks white or near-black by the fill's own luminance
    instead of assuming a color is always dark enough for white text."""
    r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return "0B0B0B" if luminance > 0.6 else "FFFFFF"


# Layout indices in python-pptx's generic built-in template (used only if
# DEFAULT_TEMPLATE_PATH is missing and no --template was given either).
DEFAULT_LAYOUTS = {
    "title": 0, "title_content": 1, "section_header": 2, "two_content": 3,
    "comparison": 4, "title_only": 5, "blank": 6, "content_caption": 7,
    "picture_caption": 8,
}

_ICON_NAMES_LINE = ", ".join(sorted(ICON_CATALOG))
_CLOUD_ICON_NAMES_LINE = ", ".join(
    f"{provider}:{name}" for provider in ("aws", "gcp", "azure", "general")
    for name in sorted(CLOUD_ICON_CATALOG.get(provider, {}))
)

SYSTEM_PROMPT = """You are the PPT Creation Agent. You have no tool access in
this call -- the brief and reference material below are already complete.
Do not ask to open files or run scripts yourself; design the plan directly
from what's given. Given a content brief (and
optionally template/reference material), design a slide-by-slide deck plan.

Respond with ONLY a single JSON object, no prose, matching this shape:
{
  "title": "<deck title>",
  "slides": [
    {"layout": "<layout name or index from the AVAILABLE LAYOUTS list>",
     "title": "...", "bullets": ["...", "..."], "notes": "<speaker notes>",
     "image_ref": null,
     "chart": {"kind": "column" | "bar" | "line" | "pie" | "area",
               "categories": ["<x-axis/category label>", "..."],
               "series": [{"name": "<series name>", "values": [<number>, "..."]}]} | null,
     "table": {"headers": ["<col>", "..."], "rows": [["<cell>", "..."], "..."]} | null,
     "chevron": {"steps": ["<short stage label>", "..."]} | null,
     "gantt": {"unit": "days" | "weeks" | "months",
               "tasks": [{"name": "<task>", "start": <number>, "duration": <number>}, "..."]} | null,
     "icons": {"items": [{"icon": "<name from the ICON CATALOG below>",
               "label": "<short headline>", "description": "<1-2 sentences, or null>"}, "..."]} | null}
  ]
}
Rules:
- Only use layout names/indices from the AVAILABLE LAYOUTS list you're given.
- Keep bullets concise (under ~15 words each), 3-6 bullets per content slide.
- Use a section_header/title_only layout to break the deck into logical
  sections if it covers multiple topics.
- Set "image_ref" to an index from AVAILABLE IMAGES only if that specific
  image is genuinely relevant to that slide; otherwise null.

Native PowerPoint objects -- use these instead of describing something as
bullets when the brief actually calls for it. At most ONE of
chart/table/chevron/gantt/icons per slide (never combine two on the same
slide), and leave "bullets" empty or very short (a one-line intro at most)
on a slide that carries one of these -- the object itself is the content,
don't duplicate it as bullets too. Leave all five null on an ordinary
bullet slide.
- "chart": a real, editable PowerPoint chart -- use for any numeric
  comparison/trend the brief asks to show (revenue, growth, breakdown by
  category, etc.). "categories" are the x-axis/slice labels; each series in
  "series" must have exactly as many "values" as there are "categories".
  Pick "kind" for what the data actually shows: "column"/"bar" for
  comparing discrete categories, "line"/"area" for a trend over time,
  "pie" only for a single categorical breakdown with a handful of slices.
  When choosing "layout" for a chart slide, prefer an AVAILABLE LAYOUTS
  entry whose "placeholder_types" shows exactly ONE "CHART" entry (e.g. a
  layout literally named "...x1...") -- the chart renders into that one
  dedicated placeholder. This schema only ever gives you one chart per
  slide, so avoid a layout whose "placeholder_types" shows CHART: 2 or more
  (a multi-chart grid layout, e.g. "2col."/"4col."/plain "charts") -- your
  single chart would only fill one of several small slots and look
  cramped; only pick one of those if you deliberately want that layout's
  other chart slots to stay empty.
- "table": a real PowerPoint table -- use for side-by-side structured
  comparison (options, specs, a matrix) that reads better as a grid than as
  bullets. Keep it readable: at most ~6 columns and ~8 data rows.
- "chevron": a horizontal row of connected arrow shapes -- use for a
  sequential process, journey, or roadmap the brief describes as steps/
  phases (3-6 short stage labels, a few words each, not full sentences).
  Note: this is a row of real AutoShapes, not a PowerPoint SmartArt graphic
  (no tool can create actual SmartArt programmatically) -- still fully
  native and editable, just visually simpler than SmartArt's own gallery.
- "gantt": a project/schedule timeline -- use when the brief describes a
  plan with tasks over time (a project plan, a rollout schedule). "start"
  and "duration" are both plain numbers in whatever "unit" you choose
  (e.g. start=2, duration=3, unit="weeks" means that task runs weeks 2-5).
  Keep it to ~8 tasks or fewer so labels stay readable. Note: PowerPoint has
  no native Gantt chart object -- this renders as a row of real, positioned
  bar shapes on a computed timeline, the standard technique for this.
- "icons": a short row of icon+label call-outs -- use when the brief
  describes a small set of distinct benefits/features/values (NOT a
  sequence -- that's "chevron") that each deserve their own icon, e.g.
  "3 reasons to choose us" or "6 core values." "icon" MUST be exactly one
  of these names (never invent one): """ + _ICON_NAMES_LINE + """.
  "label" is a short headline; "description" is an optional 1-2 sentence
  elaboration, or null. Provide exactly 3 or 6 items to match this
  template's own "3x icons Light"/"6x icons Light" grid layouts, and
  prefer one of those two layouts by name when using icons -- their
  placeholder grid is built for exactly this. These icons are plain line
  art with a fixed dark stroke baked into the image, so always pick the
  "...Light" icon layout, never a "...Dark" one, or the icon itself won't
  be visible. If the brief is about a specific cloud provider's services
  or architecture (AWS, Google Cloud, or Azure), use that provider's real
  service icons instead, named "<provider>:<service>" (e.g. "aws:ec2",
  "gcp:bigquery", "azure:sql-database"). If the brief is about automotive/
  dealership business roles instead (dealerships, cars, owners, suppliers,
  OEMs), use "general:<name>" (e.g. "general:dealership", "general:oem").
  Either way, "icon" MUST be exactly one of these (never invent one): """ + _CLOUD_ICON_NAMES_LINE + """. Never mix
  these catalogs with this template's own generic icons in the same set of
  items, and never mix two different cloud providers in the same set
  ("general" may be mixed with at most one cloud provider's set, e.g. an
  OEM's cloud-hosted supplier portal).
- For chart/table/chevron/gantt slides, prefer a spacious AVAILABLE LAYOUTS
  entry ("Title Only"/"Title and Content" style) over one with several
  fixed picture/icon placeholders that would visually compete with it --
  unless it's a chart, in which case prefer a "chart"-named layout instead
  as noted above."""


def inspect_pptx_template(path: Path) -> tuple[dict, list]:
    """A rich template (the bundled MobilityGlobal one has 46 layouts, some
    with dozens of placeholders each) makes a full idx+name dump per
    placeholder far too large -- serialized as JSON that's passed as a
    literal CLI argument to the `claude` subprocess, it blew straight past
    Windows' ~32K command-line length limit (WinError 206) well before
    even reaching the brief. A per-layout COUNT of placeholder types is a
    tiny fraction of the size and tells the LLM everything it actually
    needs to pick a layout (does it have a CHART placeholder? how many
    BODY/PICTURE slots?) without naming each one individually."""
    prs = Presentation(str(path))
    layouts = []
    for i, layout in enumerate(prs.slide_layouts):
        type_counts: dict[str, int] = {}
        for ph in layout.placeholders:
            key = ph.placeholder_format.type.name if ph.placeholder_format.type is not None else "OTHER"
            type_counts[key] = type_counts.get(key, 0) + 1
        layouts.append({"index": i, "name": layout.name, "placeholder_types": type_counts})
    return {"available_layouts": layouts}, []


def inspect_pdf_template(path: Path, images_dir: Path) -> tuple[dict, list]:
    doc = fitz.open(path)
    images_dir.mkdir(exist_ok=True, parents=True)
    pages, extracted_images = [], []
    for pno, page in enumerate(doc):
        text = page.get_text().strip()[:2000]
        page_image_indices = []
        for img in page.get_images(full=True):
            xref = img[0]
            base = doc.extract_image(xref)
            img_path = images_dir / f"page{pno + 1}_img{len(extracted_images)}.{base['ext']}"
            img_path.write_bytes(base["image"])
            page_image_indices.append(len(extracted_images))
            extracted_images.append(str(img_path))
        pages.append({"page": pno + 1, "text": text, "image_indices": page_image_indices})
    layouts = [{"index": i, "name": name} for name, i in DEFAULT_LAYOUTS.items()]
    return {"reference_pdf_pages": pages, "available_layouts": layouts,
            "note": "This PDF has no editable slide structure — treat its text/images as reference only."}, extracted_images


SLIDE_MARGIN = Inches(0.4)


def _style_chart_text(chart, slide, multi_series: bool) -> None:
    """insert_chart/add_chart leave every text element (axis tick labels,
    legend) at PowerPoint's automatic default color -- always dark,
    regardless of the slide's own background. Fine on a "...Light" layout,
    invisible on a "...Dark" one (which this template uses for real, not a
    hypothetical). This template names its light/dark variants explicitly
    (e.g. "Title and chart x1 Dark") -- use that to pick a readable color
    rather than leaving every chart's text dark-on-dark.

    Also enables the legend for multi-series charts and pies specifically
    -- pie in particular has no axis at all, so with no legend and no data
    labels (neither enabled anywhere in this file) there was previously
    nothing on the slide identifying which slice/series was which. Note
    line/area charts default to has_legend=True on their OWN already
    (unlike column/bar/pie, which default to False) -- so the color is
    applied based on whatever has_legend ends up being, not just the cases
    this function itself turns on, or that pre-existing default legend's
    text would be left at its own unset/dark default too."""
    is_dark = "dark" in (slide.slide_layout.name or "").lower()
    color = RGBColor.from_string("FFFFFF" if is_dark else "0B0B0B")

    try:
        for axis in (chart.category_axis, chart.value_axis):
            axis.tick_labels.font.color.rgb = color
    except ValueError:
        pass  # pie charts (and similar) have no axes at all

    if multi_series or chart.chart_type == XL_CHART_TYPE.PIE:
        chart.has_legend = True
        chart.legend.include_in_layout = False
    if chart.has_legend:
        chart.legend.font.color.rgb = color


def _add_chart_to_slide(slide, chart_spec: dict, slide_width: int) -> None:
    """Native, editable PowerPoint chart. Prefers inserting into the
    slide's own chart placeholder (the 6 chart-flavored layouts in
    DEFAULT_TEMPLATE_PATH all have one) so it inherits that layout's exact
    position/size; falls back to a fixed-position chart otherwise, so this
    still works with any layout or with no template at all."""
    categories = chart_spec.get("categories") or []
    series = chart_spec.get("series") or []
    if not categories or not series:
        return
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for s in series:
        chart_data.add_series(s.get("name") or "Series", s.get("values") or [])
    xl_type = CHART_KIND_TO_XL.get(chart_spec.get("kind") or "column", XL_CHART_TYPE.COLUMN_CLUSTERED)

    chart_phs = [ph for ph in slide.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER_TYPE.CHART]
    if len(chart_phs) == 1:
        graphic_frame = chart_phs[0].insert_chart(xl_type, chart_data)
    else:
        # Either no chart placeholder at all, OR the chosen layout has
        # SEVERAL (e.g. a 3-up mini-chart-grid layout like "Title and
        # Chart") -- this schema only ever supplies one chart per slide, so
        # dropping it into just one of several small mini-chart slots would
        # look cramped/oddly-placed. Full-width placement is the safe,
        # correctly-sized choice in both cases.
        content_w = slide_width - 2 * SLIDE_MARGIN
        graphic_frame = slide.shapes.add_chart(xl_type, SLIDE_MARGIN, Inches(1.7), content_w, Inches(4.8), chart_data)

    _style_chart_text(graphic_frame.chart, slide, len(series) > 1)


def _add_table_to_slide(slide, table_spec: dict, slide_width: int) -> None:
    """Native, editable PowerPoint table. Header row (if given) is filled
    with the template's primary brand color and white bold text. Spans the
    full slide width (minus margins) rather than a fixed inch value, so it
    lines up with the title above it regardless of the template's actual
    slide size (e.g. 13.33in widescreen vs. python-pptx's narrower blank
    default) -- a fixed width left a wide dead gap on a widescreen slide."""
    headers = table_spec.get("headers") or []
    rows = table_spec.get("rows") or []
    n_cols = max([len(headers)] + [len(r) for r in rows] + [0])
    n_rows = len(rows) + (1 if headers else 0)
    if n_cols == 0 or n_rows == 0:
        return
    content_w = slide_width - 2 * SLIDE_MARGIN
    table = slide.shapes.add_table(n_rows, n_cols, SLIDE_MARGIN, Inches(1.7), content_w, Inches(0.5) * n_rows).table

    r0 = 0
    if headers:
        header_fill = THEME_ACCENTS[0]
        header_text = _readable_text_color(header_fill)
        for c, h in enumerate(headers):
            cell = table.cell(0, c)
            cell.text = str(h)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(header_fill)
            for p in cell.text_frame.paragraphs:
                p.font.bold = True
                p.font.color.rgb = RGBColor.from_string(header_text)
        r0 = 1
    # Data rows keep the table's own default (light/white) cell fill --
    # explicitly dark text here rather than leaving it to inherit whatever
    # the active table style happens to default to, so it's never at the
    # mercy of a style mismatch.
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(r0 + ri, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.color.rgb = RGBColor.from_string("0B0B0B")


def _add_chevron_to_slide(slide, chevron_spec: dict, slide_width: int) -> None:
    """A row of real, editable MSO_SHAPE.CHEVRON AutoShapes for a
    sequential process/journey/roadmap -- native PowerPoint shapes, but NOT
    a SmartArt graphic (python-pptx has no API to create SmartArt at all).
    Spans the full slide width (minus margins), same reasoning as the table
    above."""
    steps = chevron_spec.get("steps") or []
    if not steps:
        return
    n = len(steps)
    total_w = slide_width - 2 * SLIDE_MARGIN
    overlap = Inches(0.3)
    seg_w = Emu(int(total_w / n) + int(overlap))
    x, y, h = SLIDE_MARGIN, Inches(3.3), Inches(1.0)
    for i, label in enumerate(steps):
        fill_hex = THEME_ACCENTS[i % len(THEME_ACCENTS)]
        shp = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, y, seg_w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill_hex)
        shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.word_wrap = True
        tf.text = str(label)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = RGBColor.from_string(_readable_text_color(fill_hex))
        x += seg_w - overlap


def _add_gantt_to_slide(slide, gantt_spec: dict, slide_width: int) -> None:
    """A row of real, positioned ROUNDED_RECTANGLE bars on a computed
    timeline scale -- PowerPoint has no native Gantt object, so this (real
    editable shapes, not a picture) is the standard technique for one.
    Timeline spans to the actual slide's right margin, same reasoning as
    the table/chevron above."""
    tasks = gantt_spec.get("tasks") or []
    if not tasks:
        return
    max_end = max((t.get("start", 0) + t.get("duration", 1)) for t in tasks) or 1
    label_w = Inches(2.4)
    chart_x0 = SLIDE_MARGIN + label_w + Inches(0.1)
    chart_x1 = slide_width - SLIDE_MARGIN
    chart_w = chart_x1 - chart_x0
    row_h, y0 = Inches(0.55), Inches(1.7)
    for i, t in enumerate(tasks):
        y = y0 + row_h * i
        label = slide.shapes.add_textbox(SLIDE_MARGIN, y, label_w, row_h)
        tf = label.text_frame
        tf.text = str(t.get("name") or "")
        tf.paragraphs[0].font.size = Pt(12)

        start, duration = t.get("start", 0), t.get("duration", 1) or 1
        bar_x = chart_x0 + Emu(int(chart_w * (start / max_end)))
        bar_w = max(Emu(int(chart_w * (duration / max_end))), Emu(1))
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bar_x, y + Inches(0.07), bar_w, Inches(0.35))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(THEME_ACCENTS[i % len(THEME_ACCENTS)])
        bar.line.fill.background()
        bar.shadow.inherit = False


def _add_icons_fallback_row(slide, items: list, slide_width: int) -> None:
    """Used when the LLM's chosen layout has no native icon grid (or fewer
    columns than items) -- a custom full-width row, same spacing math as
    _add_chevron_to_slide, so an "icons" slide never silently does nothing."""
    n = len(items)
    total_w = slide_width - 2 * SLIDE_MARGIN
    seg_w = Emu(int(total_w / n))
    icon_size, y0 = Inches(0.6), Inches(1.7)
    x = SLIDE_MARGIN
    for item in items:
        icon_path = _resolve_icon_path(item.get("icon") or "")
        if icon_path:
            icon_x = x + Emu(int((seg_w - icon_size) / 2))
            slide.shapes.add_picture(str(icon_path), icon_x, y0, width=icon_size, height=icon_size)
        label_box = slide.shapes.add_textbox(x, y0 + icon_size + Inches(0.1), seg_w, Inches(0.4))
        tf = label_box.text_frame
        tf.word_wrap = True
        tf.text = str(item.get("label") or "")
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(13)
        p.font.bold = True
        description = item.get("description")
        if description:
            desc_box = slide.shapes.add_textbox(x, y0 + icon_size + Inches(0.5), seg_w, Inches(1.2))
            dtf = desc_box.text_frame
            dtf.word_wrap = True
            dtf.text = str(description)
            dp = dtf.paragraphs[0]
            dp.alignment = PP_ALIGN.CENTER
            dp.font.size = Pt(11)
        x += seg_w


def _add_icons_to_slide(slide, icons_spec: dict, slide_width: int) -> None:
    """Native grid path: this template's "3x/6x icons Light" layouts pair
    each PICTURE placeholder with a label BODY placeholder (short) and a
    description BODY placeholder (taller) directly below it, at the same
    left x-coordinate -- confirmed via direct placeholder geometry dump.
    The 6-icon layout stacks two such rows at the SAME three x-coordinates,
    so matching is done per-picture (nearest BODY placeholders below IT
    specifically, not every BODY placeholder that merely shares its x) to
    keep the two rows from being mixed together. Falls back to a plain row
    if the chosen layout doesn't have a matching grid for this many items."""
    items = icons_spec.get("items") or []
    if not items:
        return

    tol = Inches(0.1)
    picture_phs = [ph for ph in slide.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER_TYPE.PICTURE]
    body_phs = [ph for ph in slide.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER_TYPE.BODY]

    columns = []
    for pic_ph in picture_phs:
        siblings = sorted(
            (ph for ph in body_phs if abs(ph.left - pic_ph.left) <= tol and ph.top > pic_ph.top),
            key=lambda ph: ph.top,
        )[:2]
        if not siblings:
            continue
        siblings.sort(key=lambda ph: ph.height)
        label_ph = siblings[0]
        desc_ph = siblings[1] if len(siblings) > 1 else None
        columns.append((pic_ph.top, pic_ph.left, pic_ph, label_ph, desc_ph))
    columns.sort(key=lambda c: (c[0], c[1]))  # reading order: row, then column

    if len(columns) >= len(items):
        for (_, _, pic_ph, label_ph, desc_ph), item in zip(columns, items):
            icon_path = _resolve_icon_path(item.get("icon") or "")
            if icon_path:
                pic_ph.insert_picture(str(icon_path))
            if item.get("label"):
                label_ph.text_frame.text = str(item["label"])
            if desc_ph is not None and item.get("description"):
                desc_ph.text_frame.text = str(item["description"])
        return

    _add_icons_fallback_row(slide, items, slide_width)


def _clear_existing_slides(prs: Presentation) -> None:
    """A template loaded via Presentation(path) comes in as a FULL deck,
    example content slides and all (the org template ships 30 of its own
    style-guide slides) -- strip those before adding freshly generated
    slides, so only its layouts/theme get reused, never its example
    content. Layouts/masters are untouched; only the top-level slide list
    (and each slide's now-orphaned relationship) is cleared."""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.get(qn("r:id")))
        sldIdLst.remove(sldId)


# Placeholder types that render their OWN "Click to add text" prompt (or,
# for OBJECT, a 6-icon insert-table/chart/SmartArt/picture/video/stock-image
# grid) when left empty -- visible in Normal View, not just Slide Master
# view. Several of this template's layouts put one of these directly
# on/near a plain Text Placeholder that IS getting filled (e.g. "Title and
# Content Light" has both a Content Placeholder and a Text Placeholder at
# overlapping positions) -- if the OTHER one is left untouched, its empty
# prompt visually overlaps the real content. SUBTITLE behaves the same way
# (it's just a plain text placeholder, not an auto-populated field) and a
# Title Slide with no bullets to put in it left one showing exactly that.
# FOOTER/DATE/SLIDE_NUMBER are deliberately still excluded -- those really
# are auto-populated fields, meant to stay as the template's own.
_PROMPT_SHOWING_TYPES = {
    PP_PLACEHOLDER_TYPE.BODY, PP_PLACEHOLDER_TYPE.OBJECT, PP_PLACEHOLDER_TYPE.SUBTITLE,
    PP_PLACEHOLDER_TYPE.PICTURE, PP_PLACEHOLDER_TYPE.CHART, PP_PLACEHOLDER_TYPE.TABLE,
}


def _remove_empty_content_placeholders(slide) -> None:
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx == 0 or ph.placeholder_format.type not in _PROMPT_SHOWING_TYPES:
            continue  # title, or a type that doesn't show an intrusive empty prompt
        has_content = (
            (ph.has_text_frame and ph.text_frame.text.strip())
            or getattr(ph, "has_chart", False)
            or getattr(ph, "has_table", False)
            or getattr(ph, "image", None) is not None
        )
        if not has_content:
            ph._element.getparent().remove(ph._element)


def build_pptx(plan: dict, template_path: Path | None, out_path: Path, extracted_images: list):
    if template_path and template_path.suffix.lower() == ".pptx":
        prs = Presentation(str(template_path))
        _clear_existing_slides(prs)
        layouts_by_name = {l.name.lower(): l for l in prs.slide_layouts}

        def get_layout(name_or_idx):
            if isinstance(name_or_idx, int) or str(name_or_idx).isdigit():
                idx = int(name_or_idx)
                return prs.slide_layouts[idx] if idx < len(prs.slide_layouts) else prs.slide_layouts[0]
            return layouts_by_name.get(str(name_or_idx).lower(), prs.slide_layouts[min(1, len(prs.slide_layouts) - 1)])
    else:
        prs = Presentation()

        def get_layout(name_or_idx):
            if str(name_or_idx).isdigit():
                idx = int(name_or_idx)
            else:
                idx = DEFAULT_LAYOUTS.get(str(name_or_idx).lower(), 1)
            return prs.slide_layouts[idx]

    for slide_plan in plan.get("slides", []):
        layout = get_layout(slide_plan.get("layout", "title_content"))
        slide = prs.slides.add_slide(layout)

        if slide_plan.get("title") and slide.shapes.title is not None:
            slide.shapes.title.text = slide_plan["title"]

        bullets = slide_plan.get("bullets") or []
        # Prefer an actual BODY-type text placeholder over a generic
        # OBJECT/CHART/PICTURE one that happens to also have a text frame --
        # several layouts put both at/near the same position (e.g. "Title
        # and Content Light" has a Content Placeholder AND a Text
        # Placeholder), and bullets belong in the plain text one specifically.
        text_phs = [ph for ph in slide.placeholders if ph.placeholder_format.idx != 0 and ph.has_text_frame]
        body_ph = next((ph for ph in text_phs if ph.placeholder_format.type == PP_PLACEHOLDER_TYPE.BODY), None) \
            or (text_phs[0] if text_phs else None)
        if body_ph and bullets:
            tf = body_ph.text_frame
            tf.text = bullets[0]
            for b in bullets[1:]:
                tf.add_paragraph().text = b

        image_ref = slide_plan.get("image_ref")
        if image_ref is not None and extracted_images:
            try:
                img_path = extracted_images[int(image_ref)]
                slide.shapes.add_picture(img_path, Inches(5.5), Inches(1.5), width=Inches(3.5))
            except (IndexError, ValueError):
                pass

        if slide_plan.get("chart"):
            _add_chart_to_slide(slide, slide_plan["chart"], prs.slide_width)
        if slide_plan.get("table"):
            _add_table_to_slide(slide, slide_plan["table"], prs.slide_width)
        if slide_plan.get("chevron"):
            _add_chevron_to_slide(slide, slide_plan["chevron"], prs.slide_width)
        if slide_plan.get("gantt"):
            _add_gantt_to_slide(slide, slide_plan["gantt"], prs.slide_width)
        if slide_plan.get("icons"):
            _add_icons_to_slide(slide, slide_plan["icons"], prs.slide_width)

        if slide_plan.get("notes"):
            slide.notes_slide.notes_text_frame.text = slide_plan["notes"]

        # Last, after everything real is in place -- any placeholder still
        # untouched at this point genuinely has nothing in it.
        _remove_empty_content_placeholders(slide)

    prs.save(str(out_path))


def generate(brief: str, out_path: Path, template_path: Path | None = None,
             images_dir: Path | None = None, slides: int | None = None,
             context: str | None = None) -> dict:
    """Shared by the CLI below, the library-backed server handler (see
    server.py's library_create_item), and the Workflow ppt_creator node (see
    Workflow/workflow_server.py's _run_creator_node): inspect a template,
    plan the deck with Claude, then build it deterministically. Returns the
    plan (used as this item's library metadata).

    `context` is reference content chained in from an upstream Workflow
    node (or None for the Utility Agents path, which has no such field) --
    unrelated to `template_path`'s own layout-inspection context below,
    despite the similar name.

    No explicit template_path -- falls back to DEFAULT_TEMPLATE_PATH (the
    org's own MobilityGlobal template) so every use is on-brand by default;
    only python-pptx's generic built-in layouts if that file is missing.
    """
    if template_path is None and DEFAULT_TEMPLATE_PATH.exists():
        template_path = DEFAULT_TEMPLATE_PATH

    extracted_images: list = []
    if template_path and template_path.suffix.lower() == ".pptx":
        layout_context, extracted_images = inspect_pptx_template(template_path)
    elif template_path and template_path.suffix.lower() == ".pdf":
        layout_context, extracted_images = inspect_pdf_template(template_path, images_dir or out_path.parent / "template_images")
    else:
        layout_context = {"available_layouts": [{"index": i, "name": name} for name, i in DEFAULT_LAYOUTS.items()]}

    prompt_parts = [f"BRIEF: {brief}", f"\nAVAILABLE LAYOUTS / REFERENCE MATERIAL:\n{json.dumps(layout_context, indent=2)}"]
    if extracted_images:
        prompt_parts.append(
            f"\nAVAILABLE IMAGES (reference by index 0..{len(extracted_images) - 1}): "
            f"{len(extracted_images)} image(s) extracted from the template.")
    if slides:
        prompt_parts.append(f"\nTarget slide count: about {slides}.")
    if context:
        prompt_parts.append(f"\nREFERENCE CONTENT:\n{context}")

    result = run_claude("\n".join(prompt_parts), system_prompt=SYSTEM_PROMPT, allowed_tools=[])
    plan = extract_json(result["result"])
    build_pptx(plan, template_path, out_path, extracted_images)
    return plan


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--brief", required=True, help="What the deck should be about")
    ap.add_argument("--template", help="Optional .pptx or .pdf template/reference")
    ap.add_argument("--out")
    ap.add_argument("--slides", type=int, help="Target slide count (hint only)")
    args = ap.parse_args()

    template_path = Path(args.template).resolve() if args.template else None
    out_path = Path(args.out) if args.out else Path.cwd() / "deck.pptx"

    with tempfile.TemporaryDirectory() as tmp:
        plan = generate(args.brief, out_path, template_path, Path(tmp), args.slides)

    print(f"Wrote {len(plan.get('slides', []))} slides to {out_path}")


if __name__ == "__main__":
    main()
