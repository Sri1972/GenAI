"""
Excel Parsing Agent — scans an .xlsx workbook with openpyxl, then asks Claude
(via the Claude Code CLI, on Bedrock) to turn that raw scan into a documented
metadata file: inferred column semantics, data-quality notes, and inter-sheet
relationships. The metadata file is what downstream workflow steps query
instead of re-parsing the workbook.

Usage:
    python run.py <workbook.xlsx> [--out metadata.json] [--sample-rows 15]
"""

import argparse
import datetime
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from common.claude_cli import run_claude, extract_json

SYSTEM_PROMPT = """You are the Excel Parsing Agent. You are given a raw structural
scan of a workbook (sheet names, dimensions, header rows, column names, sample
values, merged cells, and formula presence) and must produce a documented
metadata file describing it.

You have no tool access in this call. The scan provided in the prompt is
already complete and authoritative -- do not ask to run scripts, open the
file, or verify anything yourself. Answer directly from the data given.

For each column: infer its semantic meaning from its name and sample values
(not just repeat the raw type), and flag data-quality issues you notice
(blank runs, mixed types in a column, likely primary/foreign keys, obvious
duplicates). If multiple sheets look related (shared column names/values),
describe the relationship.

Respond with ONLY a single JSON object, no prose, matching this shape:
{
  "workbook": "<filename>",
  "sheets": [
    {
      "name": "...", "row_count": 0, "column_count": 0,
      "columns": [
        {"name": "...", "inferred_type": "...", "description": "...",
         "sample_values": ["..."], "quality_notes": ["..."]}
      ],
      "relationships": ["..."]
    }
  ],
  "summary": "..."
}"""


def _infer_type(values):
    types = {type(v) for v in values if v is not None}
    if not types:
        return "empty"
    if types <= {int, float}:
        return "number"
    if datetime.datetime in types or datetime.date in types:
        return "date"
    if types == {bool}:
        return "boolean"
    return "string"


def scan_workbook(path: Path, sample_rows: int) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    wb_formulas = openpyxl.load_workbook(path, data_only=False)
    sheets = []
    for ws, ws_f in zip(wb.worksheets, wb_formulas.worksheets):
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])] if rows else []
        data_rows = rows[1:]
        columns = []
        for i, col_name in enumerate(header):
            col_values = [r[i] for r in data_rows if i < len(r)]
            columns.append({
                "name": col_name,
                "raw_type": _infer_type(col_values),
                "sample_values": [str(v) for v in col_values[:sample_rows] if v is not None][:sample_rows],
                "null_count": sum(1 for v in col_values if v is None),
            })
        has_formulas = any(
            isinstance(cell.value, str) and cell.value.startswith("=")
            for row in ws_f.iter_rows() for cell in row
        )
        sheets.append({
            "name": ws.title,
            "row_count": len(data_rows),
            "column_count": len(header),
            "columns": columns,
            "has_formulas": has_formulas,
            "merged_cells": [str(r) for r in ws.merged_cells.ranges],
        })
    return {"workbook": path.name, "sheets": sheets}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--out")
    ap.add_argument("--sample-rows", type=int, default=15)
    args = ap.parse_args()

    path = Path(args.workbook).resolve()
    if not path.exists():
        sys.exit(f"File not found: {path}")

    scan = scan_workbook(path, args.sample_rows)
    prompt = (
        "Here is the raw structural scan of an Excel workbook:\n\n"
        f"{json.dumps(scan, indent=2)}\n\n"
        "Produce the metadata JSON described in your instructions."
    )

    result = run_claude(prompt, system_prompt=SYSTEM_PROMPT)
    metadata = extract_json(result["result"])

    out_path = Path(args.out) if args.out else path.with_suffix(".metadata.json")
    out_path.write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    print(f"Wrote metadata to {out_path}")


if __name__ == "__main__":
    main()
